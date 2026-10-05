#!/usr/bin/env python3
"""Build GratiWall_Presentation.pptx: 16 slides, 16:9, Fraunces + Inter, cream/ink/terracotta."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

CREAM = RGBColor.from_string('FAF7F2')
PAPER_DEEP = RGBColor.from_string('F3EDE3')
CARD = RGBColor.from_string('FFFDFA')
INK = RGBColor.from_string('1C1917')
SOFT = RGBColor.from_string('57534E')
FAINT = RGBColor.from_string('A8A29E')
LINE = RGBColor.from_string('E7E0D4')
ACCENT = RGBColor.from_string('C2410C')
ACCENT_LIGHT = RGBColor.from_string('EA580C')  # accent on dark backgrounds
GREEN = RGBColor.from_string('3F6212')
DARK_INK_SOFT = RGBColor.from_string('C9BFB2')

SERIF = 'Fraunces'
SANS = 'Inter'

SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN = 0.9

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, 'assets')
OUT = os.path.join(HERE, '..', 'GratiWall_Presentation.pptx')

prs = Presentation()
prs.slide_width = Inches(SLIDE_W)
prs.slide_height = Inches(SLIDE_H)
BLANK = prs.slide_layouts[6]


def slide(bg=CREAM):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid()
    r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def _track(run, spc):
    run._r.get_or_add_rPr().set('spc', str(spc))


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line_spacing=1.0, space_after=0):
    """runs: list of paragraphs; each paragraph is a list of (txt, font, size, color, bold, spc)."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        for (txt, font, size, color, bold, spc) in para:
            r = p.add_run()
            r.text = txt
            r.font.name = font
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
            if spc:
                _track(r, spc)
    return box


def kicker(s, label, x=MARGIN, y=0.62, color=ACCENT):
    text(s, x, y, 8, 0.3, [[(label.upper(), SANS, 10.5, color, True, 300)]])


def title(s, t, x=MARGIN, y=0.98, size=33, color=INK, w=11.5):
    text(s, x, y, w, 0.9, [[(t, SERIF, size, color, True, 0)]])


def rule(s, x, y, w, color=ACCENT, thick=2.2):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Pt(thick))
    r.fill.solid()
    r.fill.fore_color.rgb = color
    r.line.fill.background()
    r.shadow.inherit = False
    return r


def hairline(s, x, y, w, color=LINE):
    return rule(s, x, y, w, color=color, thick=0.9)


def card(s, x, y, w, h, fill=CARD, line=LINE, radius=0.06):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.adjustments[0] = radius
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.color.rgb = line
    c.line.width = Pt(1)
    c.shadow.inherit = False
    return c


def footer(s, n, dark=False):
    col = DARK_INK_SOFT if dark else FAINT
    text(s, MARGIN, 7.06, 3, 0.3, [[('GratiWall', SANS, 9, col, False, 200)]])
    text(s, SLIDE_W - MARGIN - 1, 7.06, 1, 0.3, [[(f'{n:02d}', SANS, 9, col, False, 200)]],
         align=PP_ALIGN.RIGHT)


def arrow(s, x1, y1, x2, y2, color=FAINT, w=1.4):
    conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(w)
    conn.shadow.inherit = False
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(tail)
    return conn


def screenshot_slide(n, kick, head, caption, img, accent_note=None):
    s = slide()
    text(s, MARGIN, 1.62, 3.35, 0.3, [[(kick.upper(), SANS, 10.5, ACCENT, True, 300)]])
    text(s, MARGIN, 1.98, 3.35, 1.4, [[(head, SERIF, 26, INK, True, 0)]], line_spacing=1.05)
    rule(s, MARGIN, 3.42, 0.55)
    text(s, MARGIN, 3.66, 3.35, 2.6, [[(caption, SANS, 12, SOFT, False, 0)]], line_spacing=1.3)
    if accent_note:
        text(s, MARGIN, 5.9, 3.35, 0.6, [[(accent_note, SANS, 10, GREEN, True, 60)]], line_spacing=1.2)
    # framed screenshot: soft backing plate, picture, 1px border
    w, h = 7.95, 4.97
    x, y = SLIDE_W - MARGIN - w, 1.5
    back = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.07), Inches(y + 0.09), Inches(w), Inches(h))
    back.fill.solid()
    back.fill.fore_color.rgb = LINE
    back.line.fill.background()
    back.shadow.inherit = False
    pic = s.shapes.add_picture(os.path.join(ASSETS, img), Inches(x), Inches(y), Inches(w), Inches(h))
    pic.line.color.rgb = LINE
    pic.line.width = Pt(1)
    pic.shadow.inherit = False
    footer(s, n)
    return s


# ---------------------------------------------------------------- slide 1: title
s = slide()
text(s, MARGIN, 0.95, 10, 0.3, [[('FULL STACK DEVELOPMENT · PBL 11', SANS, 11, ACCENT, True, 320)]])
text(s, MARGIN - 0.04, 1.55, 11.5, 1.6, [[('GratiWall', SERIF, 68, INK, True, 0)]])
rule(s, MARGIN, 3.24, 1.15, thick=3)
text(s, MARGIN, 3.5, 11, 0.5, [[('Campus Appreciation & Recognition Platform', SANS, 18, SOFT, False, 0)]])
text(s, MARGIN, 4.02, 11, 0.4,
     [[('Every thank-you, read by a moderator, published to a live campus wall.', SANS, 12.5, FAINT, False, 0)]])
hairline(s, MARGIN, 5.62, SLIDE_W - 2 * MARGIN)
text(s, MARGIN, 5.86, 8, 0.3, [[('TEAM', SANS, 9.5, FAINT, True, 280)]])
text(s, MARGIN, 6.14, 8, 0.35, [[('Team: [add your names]', SANS, 13, INK, False, 0)]])
text(s, SLIDE_W - MARGIN - 4.6, 5.86, 4.6, 0.3, [[('COURSE', SANS, 9.5, FAINT, True, 280)]], align=PP_ALIGN.RIGHT)
text(s, SLIDE_W - MARGIN - 4.6, 6.14, 4.6, 0.35,
     [[('24TU05MJC1 · Woxsen University', SANS, 13, INK, False, 0)]], align=PP_ALIGN.RIGHT)

# ---------------------------------------------------------------- slide 2: problem
s = slide()
kicker(s, 'The problem')
title(s, 'Good work happens quietly.')
rows = [
    ('01', 'Someone stays back to explain. No one sees it.'),
    ('02', 'Staff keep campus running. Thanks rarely reach them.'),
    ('03', 'Gratitude fades by morning. Nothing is recorded.'),
]
y = 2.15
for num, line in rows:
    text(s, MARGIN, y, 0.9, 0.6, [[(num, SERIF, 22, ACCENT, True, 0)]])
    text(s, MARGIN + 0.95, y + 0.05, 10.4, 0.6, [[(line, SERIF, 21, INK, False, 0)]])
    y += 0.92
    if num != '03':
        hairline(s, MARGIN + 0.95, y - 0.16, 10.4)
text(s, MARGIN, 5.6, 11.5, 0.8,
     [[('Appreciation exists on every campus. ', SANS, 13, SOFT, False, 0),
       ('It just has nowhere to live.', SANS, 13, ACCENT, True, 0)]], line_spacing=1.3)
footer(s, 2)

# ---------------------------------------------------------------- slide 3: objectives
s = slide()
kicker(s, 'Objectives')
title(s, 'What GratiWall does')
items = [
    ('01', 'Send', 'Students, faculty and staff write short public thank-you notes.'),
    ('02', 'Moderate', 'Every note is approved, rejected with a reason, or flagged by an admin.'),
    ('03', 'Publish live', 'Approved notes land on a big-screen wall the moment they clear review.'),
    ('04', 'Measure', 'Aggregation pipelines turn gratitude into trends, leaders and totals.'),
]
y = 2.1
for num, head, sub in items:
    text(s, MARGIN, y, 0.85, 0.5, [[(num, SERIF, 20, ACCENT, True, 0)]])
    text(s, MARGIN + 0.9, y - 0.02, 3.1, 0.5, [[(head, SERIF, 19, INK, True, 0)]])
    text(s, 4.6, y + 0.04, 7.8, 0.5, [[(sub, SANS, 13, SOFT, False, 0)]])
    y += 1.06
    if num != '04':
        hairline(s, MARGIN, y - 0.24, SLIDE_W - 2 * MARGIN)
footer(s, 3)

# ---------------------------------------------------------------- slide 4: roles
s = slide()
kicker(s, 'Users & roles')
title(s, 'Four roles, one wall')
roles = [
    ('Student', 'Sends notes to anyone on campus, posts anonymously, tracks moderation status.'),
    ('Faculty', 'Thanks students and peers publicly; receives notes by email when published.'),
    ('Staff', 'Mess, library, admin and hostel staff finally get named and thanked.'),
    ('Admin', 'Reads the queue, approves or rejects with reasons, watches the analytics.'),
]
cw, gap = 2.79, 0.2
x = MARGIN
for name, desc in roles:
    card(s, x, 2.15, cw, 3.5)
    rule(s, x + 0.28, 2.5, 0.42, thick=2.4)
    text(s, x + 0.28, 2.72, cw - 0.56, 0.5, [[(name, SERIF, 19, INK, True, 0)]])
    text(s, x + 0.28, 3.34, cw - 0.56, 2.1, [[(desc, SANS, 12, SOFT, False, 0)]], line_spacing=1.3)
    x += cw + gap
text(s, MARGIN, 6.0, 11.5, 0.5,
     [[('Roles are enforced twice: ', SANS, 12.5, SOFT, False, 0),
       ('route guards in React, JWT middleware in Express.', SANS, 12.5, INK, True, 0)]])
footer(s, 4)

# ---------------------------------------------------------------- slide 5: architecture
s = slide()
kicker(s, 'System architecture')
title(s, 'Simple, live, observable')

def node(x, y, w, h, head, sub, dashed=False):
    c = card(s, x, y, w, h)
    if dashed:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn('a:prstDash'), {'val': 'dash'}))
    text(s, x + 0.24, y + 0.18, w - 0.48, 0.4, [[(head, SERIF, 15.5, INK, True, 0)]])
    text(s, x + 0.24, y + 0.58, w - 0.48, h - 0.7, [[(sub, SANS, 10.5, SOFT, False, 0)]], line_spacing=1.2)
    return c

node(0.9, 2.5, 3.1, 1.25, 'React client', 'Vite, React Router, axios. Public wall and admin views.')
node(5.12, 2.5, 3.1, 1.25, 'Express API', 'JWT auth, role guard, Mongoose models, REST endpoints.')
node(9.33, 2.5, 3.1, 1.25, 'MongoDB', 'Users and Notes collections. Aggregation pipelines.')
arrow(s, 4.02, 3.12, 5.1, 3.12, color=SOFT)
arrow(s, 8.24, 3.12, 9.31, 3.12, color=SOFT)
text(s, 4.02, 2.78, 1.1, 0.25, [[('REST · JWT', SANS, 8.5, FAINT, True, 120)]], align=PP_ALIGN.CENTER)
text(s, 8.24, 2.78, 1.1, 0.25, [[('ODM', SANS, 8.5, FAINT, True, 120)]], align=PP_ALIGN.CENTER)

node(1.7, 4.85, 4.4, 1.05, 'Socket.IO', 'note:published and applause events pushed to every open wall.', dashed=True)
node(7.25, 4.85, 4.4, 1.05, 'Nodemailer', 'Recipient notified by email on publish; console fallback without SMTP.', dashed=True)
arrow(s, 5.62, 3.77, 4.4, 4.83, color=ACCENT)
arrow(s, 7.2, 3.77, 8.9, 4.83, color=ACCENT)
text(s, MARGIN, 6.35, 11.5, 0.4,
     [[('One process serves API and sockets. If MongoDB is down the API degrades gracefully with a clear 503.', SANS, 11.5, FAINT, False, 0)]])
footer(s, 5)

# ---------------------------------------------------------------- slide 6: tech stack
s = slide()
kicker(s, 'Tech stack')
title(s, 'MERN, plus the pieces that make it live')
stack = [
    ('React 18', 'UI'), ('Vite', 'BUILD'), ('React Router', 'ROUTING'),
    ('Express', 'API'), ('MongoDB', 'DATABASE'), ('Mongoose', 'ODM'),
    ('Socket.IO', 'REALTIME'), ('JWT + bcrypt', 'AUTH'), ('Recharts', 'ANALYTICS'),
]
cw, ch, gap = 3.71, 1.15, 0.2
for i, (name, tag) in enumerate(stack):
    x = MARGIN + (i % 3) * (cw + gap)
    y = 2.15 + (i // 3) * (ch + gap)
    card(s, x, y, cw, ch)
    text(s, x + 0.26, y + 0.2, cw - 0.52, 0.3, [[(tag, SANS, 8.5, ACCENT, True, 260)]])
    text(s, x + 0.26, y + 0.52, cw - 0.52, 0.45, [[(name, SERIF, 17, INK, True, 0)]])
text(s, MARGIN, 6.35, 11.5, 0.4,
     [[('Monorepo with npm workspaces. One npm run dev starts client and server together.', SANS, 11.5, FAINT, False, 0)]])
footer(s, 6)

# ---------------------------------------------------------------- slide 7: data model
s = slide()
kicker(s, 'Data model')
title(s, 'Two collections')

def schema_card(x, name, fields):
    card(s, x, 2.05, 5.66, 4.35)
    text(s, x + 0.3, 2.28, 3, 0.3, [[('COLLECTION', SANS, 8.5, FAINT, True, 260)]])
    text(s, x + 0.3, 2.52, 4, 0.45, [[(name, SERIF, 20, INK, True, 0)]])
    y = 3.14
    for fname, ftype in fields:
        text(s, x + 0.3, y, 2.6, 0.3, [[(fname, SANS, 11.5, INK, True, 0)]])
        text(s, x + 2.7, y, 2.75, 0.3, [[(ftype, SANS, 11, SOFT, False, 0)]])
        y += 0.335

schema_card(MARGIN, 'User', [
    ('name', 'String'),
    ('email', 'String, unique'),
    ('passwordHash', 'String, bcrypt'),
    ('role', 'student · faculty · staff · admin'),
    ('department', 'one of five schools'),
    ('timestamps', 'createdAt, updatedAt'),
])
schema_card(MARGIN + 5.66 + 0.22, 'Note', [
    ('sender', 'ref User, always stored'),
    ('senderAnonymous', 'Boolean, hidden on wall'),
    ('recipientName + recipient', 'String + optional ref User'),
    ('category', 'one of four kinds'),
    ('message', 'String, max 500 chars'),
    ('status', 'pending · approved · rejected · flagged'),
    ('department', 'denormalized for filters'),
    ('applause', 'Number, live counter'),
    ('moderatedBy / At / reason', 'ref User, Date, String'),
])
footer(s, 7)

# ---------------------------------------------------------------- slide 8: lifecycle
s = slide()
kicker(s, 'Note lifecycle')
title(s, 'From draft to the wall in one review')
steps = [
    ('Submitted', 'author writes up to 500 characters'),
    ('Pending', 'enters the moderation queue'),
    ('Moderation', 'approve · reject · flag'),
    ('Published', 'pushed live over Socket.IO'),
    ('Notified', 'email if recipient is registered'),
]
bw, bh, gap = 2.1, 1.35, 0.32
total = 5 * bw + 4 * gap
x = (SLIDE_W - total) / 2
y = 2.75
for i, (head, sub) in enumerate(steps):
    accent_step = i == 2
    c = card(s, x, y, bw, bh, fill=(CARD if not accent_step else INK),
             line=(LINE if not accent_step else INK))
    text(s, x + 0.18, y + 0.2, bw - 0.36, 0.35,
         [[(head, SERIF, 15, INK if not accent_step else CREAM, True, 0)]])
    text(s, x + 0.18, y + 0.6, bw - 0.36, 0.7,
         [[(sub, SANS, 10, SOFT if not accent_step else DARK_INK_SOFT, False, 0)]], line_spacing=1.2)
    if i < 4:
        arrow(s, x + bw + 0.03, y + bh / 2, x + bw + gap - 0.03, y + bh / 2, color=ACCENT, w=1.6)
    x += bw + gap
text(s, MARGIN, 5.15, 11.53, 0.9,
     [[('Rejections carry a reason back to the sender. Flags park a note for a second look. ', SANS, 12.5, SOFT, False, 0)],
      [('Nothing reaches the wall unread.', SANS, 12.5, ACCENT, True, 0)]],
     align=PP_ALIGN.CENTER, line_spacing=1.35)
footer(s, 8)

# ---------------------------------------------------------------- slides 9-13: screenshots
screenshot_slide(9, 'The wall', 'Live, public, always on',
                 'Approved notes appear on every connected screen the moment they clear review. A spotlight rotates a featured note every 2.5 seconds; filters slice the wall by kind of note and department.',
                 'wall-light.png', accent_note='DESIGNED FOR CAFETERIA SCREENS')
screenshot_slide(10, 'Dark mode', 'One toggle, full theme',
                 'The wall respects the OS preference on first visit and remembers the choice. Every surface, chart, badge and form adapts; contrast stays readable on projectors at night.',
                 'wall-dark.png', accent_note='PERSISTS VIA LOCALSTORAGE')
screenshot_slide(11, 'Submit & anonymity', 'Thirty seconds to thank someone',
                 'A searchable picker finds any registered member of campus, or a name can be typed freehand. Anonymous notes hide the sender on the wall; moderators still see who wrote them.',
                 'submit.png', accent_note='500 CHARACTERS, FOUR CATEGORIES')
screenshot_slide(12, 'Moderation', 'Every note is read first',
                 'The queue shows sender, recipient, category and full message with time ago. Approve publishes instantly and emails the recipient; reject sends a reason back to the author; flag parks it.',
                 'admin-queue.png', accent_note='ADMIN SEES REAL SENDERS')
screenshot_slide(13, 'Analytics', 'The culture, measured',
                 'MongoDB aggregation pipelines power the dashboard: notes per department, weekly trends for each category over eight weeks, the most appreciated people, and totals by status.',
                 'analytics.png', accent_note='FOUR PIPELINES, ONE ENDPOINT')

# ---------------------------------------------------------------- slide 14: beyond the brief
s = slide()
kicker(s, 'Beyond the brief')
title(s, 'Details that make it feel finished')
extras = [
    ('TV mode', 'Fullscreen view with enlarged type for cafeteria displays; Esc exits.'),
    ('Wall filters', 'Filter by category and department, server-side, with pagination.'),
    ('Applause', 'A clap counter on every note, one per browser, live over sockets.'),
    ('Recipient pages', 'Every name on the wall links to a page of all their notes.'),
    ('Intro animation', 'An editorial title card on each full load; skipped for reduced motion.'),
    ('Dark mode', 'A complete second palette, OS-aware, remembered per browser.'),
]
cw, ch, gap = 3.71, 1.55, 0.2
for i, (head, sub) in enumerate(extras):
    x = MARGIN + (i % 3) * (cw + gap)
    y = 2.1 + (i // 3) * (ch + gap)
    card(s, x, y, cw, ch)
    rule(s, x + 0.26, y + 0.3, 0.34, thick=2.2)
    text(s, x + 0.26, y + 0.48, cw - 0.52, 0.35, [[(head, SERIF, 15.5, INK, True, 0)]])
    text(s, x + 0.26, y + 0.86, cw - 0.52, 0.62, [[(sub, SANS, 10.5, SOFT, False, 0)]], line_spacing=1.2)
footer(s, 14)

# ---------------------------------------------------------------- slide 15: evaluation mapping
s = slide()
kicker(s, 'Evaluation mapping')
title(s, 'Where each component is demonstrated')
mapping = [
    ('Problem Definition & Objectives', '5', 'The wall itself: gratitude made visible, moderated, measurable.'),
    ('Design & Methodology', '5', 'Architecture, data model and lifecycle diagrams in this deck.'),
    ('Implementation & Analysis', '8', 'A running MERN app: live demo plus the analytics dashboard.'),
    ('Project Report', '5', 'README with setup, API reference, seed data and credentials.'),
    ('Presentation & Viva', '2', 'This deck, the live wall, and a walkthrough of the queue.'),
]
y = 2.2
text(s, MARGIN, y, 5.2, 0.3, [[('COMPONENT', SANS, 9, FAINT, True, 260)]])
text(s, 6.6, y, 0.9, 0.3, [[('MARKS', SANS, 9, FAINT, True, 260)]])
text(s, 7.7, y, 4.7, 0.3, [[('EVIDENCE', SANS, 9, FAINT, True, 260)]])
y += 0.42
hairline(s, MARGIN, y, SLIDE_W - 2 * MARGIN)
y += 0.14
for comp, marks, where in mapping:
    text(s, MARGIN, y, 5.2, 0.4, [[(comp, SANS, 13, INK, True, 0)]])
    text(s, 6.6, y, 0.9, 0.4, [[(marks, SERIF, 14, ACCENT, True, 0)]])
    text(s, 7.7, y + 0.02, 4.75, 0.55, [[(where, SANS, 11.5, SOFT, False, 0)]], line_spacing=1.15)
    y += 0.78
    hairline(s, MARGIN, y - 0.12, SLIDE_W - 2 * MARGIN)
footer(s, 15)

# ---------------------------------------------------------------- slide 16: closing
s = slide(bg=INK)
text(s, MARGIN, 1.15, 10, 0.3, [[('GRATIWALL', SANS, 11, ACCENT_LIGHT, True, 340)]])
text(s, MARGIN - 0.04, 2.3, 11.5, 1.5, [[('Thank you.', SERIF, 58, CREAM, True, 0)]])
rule(s, MARGIN, 3.85, 1.15, color=ACCENT_LIGHT, thick=3)
text(s, MARGIN, 4.15, 11, 0.5, [[('Questions?', SANS, 17, DARK_INK_SOFT, False, 0)]])
text(s, MARGIN, 6.35, 11.5, 0.4,
     [[('24TU05MJC1 · Full Stack Development · Woxsen University', SANS, 10.5, RGBColor.from_string('8F8478'), False, 140)]])
footer(s, 16, dark=True)

prs.save(OUT)
print(f'saved {os.path.abspath(OUT)} with {len(prs.slides.slides if hasattr(prs.slides, "slides") else prs.slides._sldIdLst)} slides')
